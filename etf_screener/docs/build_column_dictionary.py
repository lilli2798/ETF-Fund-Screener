"""Build docs/report_column_dictionary.xlsx: what every report column means.

One row per column name; columns that only differ by period / variant
(e.g. Maximum Drawdown (1Y/3Y/5Y/10Y)) are listed once with their variants.
Each row has the definition, formula, optimal value, conditions / caveats,
source, and how much it counts in the profile score (read from the
profile YAML, so re-run this after changing weights).

Usage (from etf_screener/):
    python docs/build_column_dictionary.py [profile.yaml] [report.xlsx]
The profile defaults to input_files/profile_long_term.yaml, which writes
docs/report_column_dictionary.xlsx; any other profile writes
docs/report_column_dictionary_<key>.xlsx. Without a report, the newest
report with that profile's columns (normally the combined
results_profile_LT_ST_*.xlsx) in the YAML's out_path is used to get the
list of report columns. The "Used In" column and concept weights follow
the chosen profile; the other profiles' Score / Rank columns get
their own rows with their own YAML weights.
"""
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

PROJECT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT))
sys.path.append(str(PROJECT / "yahoo_returns"))
from config import PURCHASE_RULES, VOLATILE_CONFIG_PATH  # noqa: E402
from input_file import discover_profiles, load_profile_input  # noqa: E402
from volatile_periods import load_periods  # noqa: E402

ALL_PROFILES = discover_profiles()  # every enabled input_files/profile_*.yaml
DEFAULT_PROFILE = Path(ALL_PROFILES[0].path)  # the report's main profile (LT)
args = sys.argv[1:]
yaml_args = [a for a in args if a.endswith((".yaml", ".yml"))]
report_args = [a for a in args if a.endswith(".xlsx")]

PROFILE_PATH = Path(yaml_args[0]) if yaml_args else DEFAULT_PROFILE
profile = load_profile_input(str(PROFILE_PATH))
KEY = profile.profile_name
PA = f"Profile {KEY}"  # e.g. "Profile LT", as in the report's column names
OUT = PROJECT / ("docs/report_column_dictionary.xlsx" if PROFILE_PATH.resolve() == DEFAULT_PROFILE.resolve()
                 else f"docs/report_column_dictionary_{KEY}.xlsx")
if report_args:
    REPORT = Path(report_args[0])
else:
    # newest report with this profile's columns: results_profile_LT_ST_... or results_profile_LT_...
    reports = sorted((PROJECT / profile.out_path).glob(f"results_profile_*{KEY}_*.xlsx"), key=lambda r: r.stat().st_mtime)
    reports = [r for r in reports if not r.name.startswith("~$")]
    if not reports:
        sys.exit(f"No results_profile_*{KEY}_*.xlsx in {PROJECT / profile.out_path} - run the screener first or pass a report path.")
    REPORT = reports[-1]
print(f"Profile {KEY} ({PROFILE_PATH.name}); report columns from: {REPORT.name}")
W = profile.thresholds["weights"]
CW = profile.thresholds["concept_weights"]

def rule(key: str, fmt) -> str:
    """A shared purchase rule from config.PURCHASE_RULES, e.g. 'max_front_load = 0%'."""
    v = PURCHASE_RULES.get(key)
    return f"{key} = {'no limit' if v is None else fmt(v)}"


def pct(v):
    return f"{v:g}%"


def dollars(v):
    return f"${v:,.0f}"

MS, YH, CALC, ET = "Morningstar", "Yahoo Finance (calculated by screener)", "Screener (calculated)", "eTrade / Morningstar export"
CONCEPT_LABEL = {
    "long_term_return_performance": "Long-term Return", "short_term_return_performance": "Short-term Return",
    "risk_adjusted": "Risk-adjusted", "volatility": "Volatility", "tracking": "Tracking",
    "liquidity_size": "Liquidity/Size", "quality_valuation": "Quality/Valuation", "costs": "Costs", "tax_income": "Tax/Income",
    "resilience": "Resilience",
}
WITHIN = ("Scored as a percentile rank WITHIN the fund's Morningstar Category, so a fund is only compared with "
          "funds of the same type.")
NORM = ("Each input is turned into a 0-100 percentile rank within the Morningstar Category "
        "(score = rank% x 100; for 'lower is better' inputs score = 100 - rank% x 100), then "
        "score = Σ(weight x input score) / Σ(weights of inputs the fund has).")


# Inputs that are never in the data, so the code never gives them weight
ABSENT = {("liquidity_size", "trading_volume"), ("tax_income", "dividend_per_share_forward"),
          ("tax_income", "dividend_per_share_1y_avg"), ("tax_income", "dividend_per_share_3y_avg")}
TOTAL_CW = sum(v for v in W.values() if v)


def concept_share(concept):
    """Effective share of the profile score (the code divides by the sum of weights)."""
    return W.get(concept, 0) / TOTAL_CW


def inner_share(concept, key):
    inner = CW.get(concept, {})
    total = sum(v for k, v in inner.items() if (concept, k) not in ABSENT)
    return inner[key] / total if total else 0.0


def used(concept, keys):
    """Score usage text for one or more (variant label, weight key) pairs."""
    cw = W.get(concept, 0)
    parts = []
    for label, key in keys:
        inner = CW.get(concept, {}).get(key)
        if not inner or (concept, key) in ABSENT:
            continue
        prefix = f"{label}: " if label else ""
        share = inner_share(concept, key)
        parts.append(f"{prefix}YAML weight {inner:g} = {share:.0%} of {CONCEPT_LABEL[concept]} (≈{concept_share(concept) * share * 100:.1f}% of {PA})")
    if not parts:
        return "No"
    if cw == 0:
        return f"No - {CONCEPT_LABEL[concept]} has 0 weight in {PA} (inputs: " + "; ".join(parts) + ")"
    return "Yes - " + "; ".join(parts) + f". {CONCEPT_LABEL[concept]} = {concept_share(concept):.1%} of {PA}"


# --------------------------------------------------------------------------
# Group report columns by base name
# --------------------------------------------------------------------------
report_cols = list(pd.read_excel(REPORT, nrows=0).columns)
extra_cols = ["Missing From Yahoo"]


def base_of(col):
    m = re.match(r"^(Volatile (?:Decline|Recovery)): (.+)$", col)
    if m:
        return m.group(1), m.group(2)
    m = re.match(r"^((?:Yearly|Volatile) )(Top \d+%) Count$", col)
    if m:
        return f"{m.group(1)}Top X% Count", m.group(2)
    m = re.match(r"^(Sharpe|Z Score) (\dy)$", col)
    if m:
        return m.group(1), m.group(2).upper()
    m = re.match(r"^(.+?) \((.+)\)$", col)
    if m and m.group(1) != "Equity Style Box":
        return m.group(1), m.group(2).replace(" Monthly", "")
    return col, ""


groups = {}
for col in report_cols + extra_cols:
    b, v = base_of(col)
    groups.setdefault(b, [])
    groups[b].append(v or "current")

PERIODS = {p["name"]: p for p in load_periods(Path(VOLATILE_CONFIG_PATH))}

# --------------------------------------------------------------------------
# Definitions: base name -> (definition, formula/notes, optimal, condition, source, score)
# --------------------------------------------------------------------------
D = {}

D["Ticker"] = ("Trading symbol of the ETF or mutual fund share class.", "", "n/a (identifier)",
               "Mutual fund tickers have 5 letters ending in X (e.g. VFIAX); ETFs usually 3-4 letters. Rows starting with ^ at the bottom of the overview are benchmark indexes (S&P 500, Dow, Nasdaq).", ET, "No")
D["Name"] = ("Full name of the fund / share class.", "", "n/a",
             "The same fund can appear several times as different share classes (Investor, Admiral, Institutional...) with different costs and minimums.", ET, "No")
D["Morningstar Category"] = ("Morningstar's peer group for the fund, based on what it actually held over the last 3 years (e.g. Large Blend, Intermediate Core Bond, Technology).", "", "n/a (grouping)",
                             "Every score and in-category rank compares a fund ONLY with funds in the same category. 132 categories in this report (a few appear twice under ETF vs mutual-fund naming, e.g. 'Mid-Cap Blend' / 'Mid Blend'). "
                             "Sector caveat: categories are coarse - e.g. 'Technology' mixes semiconductor-heavy and software-heavy funds, which behave very differently; see the Yahoo Subsector column and the 'Scoring Concepts' tab.", MS, "Used as the peer group for all scores")
D["Asset Class"] = ("Broad asset class Morningstar assigns: US Equity, International Equity, Sector Equity, Taxable Bond, Municipal Bond, Allocation, Alternative, Commodities, Money Market, ...", "", "n/a (grouping)",
                    "Compare funds within the same asset class before comparing returns.", MS, "No")
D["Primary Benchmark"] = ("Index the fund says it is measured against in its prospectus.", "", "n/a",
                          "Morningstar's Alpha, Beta, capture ratios and tracking error may use a category index instead. If the benchmark doesn't fit the fund's holdings, those measures are misleading.", MS, "No")
D["Equity Style Box (Funds)"] = ("Morningstar Style Box: company size (Large / Mid / Small) x style (Value / Blend / Growth) of the stock holdings.", "", "n/a (grouping)",
                         "Blank for bond, commodity and most alternative funds.", MS, "No")
D["Inception Date"] = ("Date the fund (or this share class) started.", "", "Older = longer track record",
                       "For ETFs converted from a mutual fund, Morningstar may show the old fund's start date while Yahoo only has prices since the ETF launch. Funds younger than 3 years have no 3Y metrics.", MS, "No")
D["Broad Category Group"] = ("Morningstar's top-level grouping of categories: Equity, Fixed Income, Allocation, Alternative Strategies, Real Assets, Hybrid Securities, Capital Preservation, Other.", "", "n/a (grouping)", "8 groups in this report.", MS, "No")
D["TTM Yield"] = ("Trailing 12-month yield: income paid over the last 12 months as a percent of the current price.", "TTM Yield = distributions of last 12 months / current price x 100",
                  "Higher = more income (not 'better' by itself)",
                  "Yields above ~10% usually come from option-income, leveraged or return-of-capital funds and may not last. Compare with SEC 30-Day Yield and within category.", MS, "No")
D["Total Return"] = ("Total return (price change + reinvested dividends/interest, after fees), in percent. 1W-1Y are cumulative; 2Y and longer are annualized (average per year). 'Since Inception' is annualized from the start date.",
                     "Cumulative: (Ending value + distributions) / Beginning value - 1.  Annualized (n years): (1 + cumulative)^(1/n) - 1",
                     "Higher is better",
                     "Compare within category - a 5% year is great for a bond fund and poor for a tech fund. Blank when the fund is younger than the period. "
                     "Sector caveat: multi-year returns are strongly driven by sub-sector timing (e.g. AI semiconductors vs software), which within-category ranking does not correct for. "
                     "'Since Inception' is not comparable across funds with different start dates. "
                     "A fund younger than a period has that return blank, and it simply drops out of the fund's score.",
                     MS, used("short_term_return_performance", [(p, f"return_{p.lower()}") for p in ["1W", "1M", "2M", "QTD", "3M", "6M", "9M", "YTD"]]) + "  |  " +
                     used("long_term_return_performance", [(p, f"return_{p.lower()}") for p in ["1Y", "2Y", "3Y", "4Y", "5Y", "6Y", "7Y", "8Y", "9Y", "10Y", "15Y", "20Y"]]))
D["Yearly Return"] = ("The fund's total return for one calendar year (Jan 1 - Dec 31), in percent, from Yahoo adjusted prices.",
                      "Yearly Return = Adj Close on last trading day of year / Adj Close on last trading day of prior year - 1",
                      "Higher is better",
                      "Compare with the S&P 500 / Dow / Nasdaq rows at the bottom of the overview. Blank if the fund didn't exist that whole year or isn't in the yearly returns file yet.", YH, "No")

stat_defs = {
    "Beat All Indexes": ("Number of {u} in which the fund's return was higher than ALL three benchmark indexes (S&P 500 ^GSPC, Dow Jones ^DJI, Nasdaq ^IXIC).",
                         "Count of {u} where fund return > max(S&P 500, Dow, Nasdaq return)", "Higher is better",
                         "Raw count: depends on history length - use '{p} Beat %' to compare funds. Bond and balanced funds rarely beat all-stock indexes in rising markets; compare within category."),
    "Trailed All Indexes": ("Number of {u} in which the fund's return was lower than ALL three benchmark indexes.",
                            "Count of {u} where fund return < min(S&P 500, Dow, Nasdaq return)", "Lower is better",
                            "Raw count; read with '{p} Periods With Data'. Defensive funds often trail stock indexes in strong up years - expected, not bad."),
    "Periods With Data": ("Number of {u} for which the fund has a return.", "Count of non-blank {u}", "Higher = longer record (more reliable stats)",
                          "Not a quality measure by itself. Treat percentages and bucket counts from funds with only 1-3 {u} as unreliable."),
    "Beat %": ("Share of the fund's {u} in which it beat all three indexes, in percent.", "Beat % = Beat All Indexes / Periods With Data x 100", "Higher is better (0-100)",
               "Comparable across funds with different history lengths, but noisy when Periods With Data is small."),
    "Trailed %": ("Share of the fund's {u} in which it trailed all three indexes, in percent.", "Trailed % = Trailed All Indexes / Periods With Data x 100", "Lower is better (0-100)",
                  "Noisy when Periods With Data is small."),
    "Beat-Trail Diff": ("Beat % minus Trailed %, in percentage points.", "Beat-Trail Diff = Beat % - Trailed %", "Higher is better (-100 to +100; > 0 = beat more often than trailed)",
                        "One-number summary of the two above; still unreliable with few periods."),
    "Top X% Count": ("Number of {u} in which the fund's return fell in a given 10% band of ALL funds in the report (benchmark indexes excluded). Top 10% = best tenth (90th-100th percentile); Top 20% = 80th-90th band; ... Top 90% = 10th-20th band.",
                     "Per {o}: percentile = rank of the fund's return among all funds with a return that {o} (average rank, pct). Count the {u} whose percentile falls in the band.",
                     "Higher in Top 10%-50% bands; lower in Top 60%-90% bands",
                     "IMPORTANT: bands are NOT cumulative - 'Top 20% Count' counts only the 80th-90th band. Add Top 10% + Top 20% to get 'in the top 20%'. Ranked against ALL funds (all categories), so stock funds dominate the top in up {u} and bond funds in down {u}."),
    "Bottom 10% Count": ("Number of {u} in which the fund's return was in the worst 10% of all funds in the report.",
                         "Count of {u} with percentile <= 10%", "Lower is better",
                         "Ranked against all funds in all categories. Many bottom-10% {u} = often among the worst of the whole universe."),
}
for prefix, u, o, src in (("Yearly", "calendar years", "year", YH), ("Volatile", "volatile periods", "period", YH)):
    for stat, (d, f, opt, c) in stat_defs.items():
        D[f"{prefix} {stat}"] = (d.format(u=u, o=o, p=prefix), f.format(u=u, o=o), opt, c.format(u=u, o=o, p=prefix), CALC, "No")
    # Inputs of the Resilience score
    D[f"{prefix} Beat %"] = D[f"{prefix} Beat %"][:5] + (used("resilience", [("", f"{prefix.lower()}_beat_pct")]),)
    D[f"{prefix} Bottom 10% Count"] = D[f"{prefix} Bottom 10% Count"][:5] + (
        used("resilience", [("As a share of Periods With Data", f"{prefix.lower()}_bottom_10_share")]),)

D["Volatile Decline"] = ("Fund's total return during a market crisis: from the S&P 500 high (start date) to its low (end date), in percent. One column per crisis - see the Variants column and the 'Volatile Periods' tab for dates.",
                         "Volatile Decline = Adj Close on low date / Adj Close on start date - 1", "Higher (smaller loss) is better",
                         "Shows downside protection. Bond / defensive funds naturally lose less - compare within category. Blank if the fund didn't exist then. Bear markets also have a Recovery column; 'decline legs' and spikes do not.", YH, "No")
D["Volatile Recovery"] = ("Fund's total return from the crisis low until the S&P 500 regained its old high (bear markets only), in percent.",
                          "Volatile Recovery = Adj Close on recovery date / Adj Close on low date - 1", "Higher is better",
                          "Read together with Volatile Decline - the ideal fund loses less AND recovers well.", YH, "No")
D["Sharpe Ratio"] = ("Morningstar's Sharpe ratio: return above the risk-free rate per unit of volatility, from monthly returns over 1, 3, 5 or 10 years.",
                     "Sharpe = (Rp - Rf) / σp, annualized; Rp = fund return, Rf = risk-free (T-bill) rate, σp = standard deviation of monthly returns",
                     "Higher is better (> 1 good, > 2 very good, < 0 = did worse than cash)",
                     "1Y is short and noisy; prefer 3Y/5Y. Bond funds have different typical levels than stock funds - compare within category. Inherits sector distortion through σp (semis more volatile than software).",
                     MS, used("risk_adjusted", [("1Y", "sharpe_1y"), ("3Y", "sharpe_3y"), ("5Y", "sharpe_5y"), ("10Y", "sharpe_10y")]))
D["Standard Deviation"] = ("Annualized standard deviation of monthly returns, in percent - how much the return swings up and down.",
                           "σ = sqrt( Σ(Ri - R̄)² / (n - 1) ) x sqrt(12), Ri = monthly return",
                           "Lower = steadier",
                           "Up and down swings count the same. Typical: bond funds 3-7%, broad stock funds 15-20%. HIGH sector distortion: semiconductor-heavy funds are far more volatile than software-heavy funds in the same category.",
                           MS, used("volatility", [("1Y", "stdev_1y"), ("3Y", "stdev_3y"), ("5Y", "stdev_5y"), ("10Y", "stdev_10y")]))
D["Worst Three Month Return"] = ("Lowest return over any rolling 3-month window in Morningstar's history window, in percent.", "", "Higher (less negative) is better",
                                 "A gut-check: could you hold the fund through this loss?", MS, used("resilience", [("", "worst_3m_return")]))
D["Best Three Month Return"] = ("Highest return over any rolling 3-month window, in percent.", "", "Higher is better, but read with the worst 3 months",
                                "A big best AND big worst = very volatile fund (often leveraged or single-sector).", MS, "No")
D["Morningstar Risk Rating"] = ("Morningstar's downside-risk rating vs the category: Low, Below Average, Average, Above Average, High.",
                                "Based on Morningstar Risk (downside variation of monthly returns) ranked within category: top 10% riskiest = High, next 22.5% = Above Average, middle 35% = Average, next 22.5% = Below Average, bottom 10% = Low",
                                "Low / Below Average for low-risk investing",
                                "Relative to the category - 'Low' risk in an emerging-markets category is still risky in absolute terms.", MS, "No")
D["Upside Capture Ratio"] = ("Share of the benchmark's gains the fund captured in months when the benchmark rose, in percent (100 = same as benchmark).",
                             "Upside Capture = (fund's compounded return in benchmark-up months / benchmark's return in those months) x 100",
                             "Higher is better",
                             "Read with Downside Capture: upside > downside is ideal. Values far beyond ±300 happen when the benchmark barely moved - ignore those. Semi-heavy funds amplify both up and down capture.",
                             MS, used("risk_adjusted", [("1Y", "upside_1y"), ("3Y", "upside"), ("5Y", "upside_5y"), ("10Y", "upside_10y")]))
D["Downside Capture Ratio"] = ("Share of the benchmark's losses the fund took in months when the benchmark fell, in percent (60 = lost only 60% as much).",
                               "Downside Capture = (fund's compounded return in benchmark-down months / benchmark's return in those months) x 100",
                               "Lower is better (< 100 = protects in bad months)",
                               "Negative = the fund rose while the benchmark fell. Extreme values are unreliable. Scored inverted.",
                               MS, used("risk_adjusted", [("1Y", "downside_1y"), ("3Y", "downside"), ("5Y", "downside_5y"), ("10Y", "downside_10y")]))
D["Maximum Drawdown"] = ("Largest peak-to-bottom loss over the period, in percent (a negative number).",
                         "Max Drawdown = min over t of ( Value_t / max(Value_s for s ≤ t) - 1 )",
                         "Closer to 0 is better",
                         "The worst pain an investor would have felt. 5Y includes the 2022 bear market. HIGH sector distortion (cyclical semis vs software). Scored as-is (not inverted): the higher, less negative value scores better.",
                         MS, used("volatility", [("1Y", "drawdown_1y"), ("3Y", "drawdown_3y"), ("5Y", "drawdown_5y"), ("10Y", "drawdown_10y")]))
D["Portfolio Risk Score"] = ("Morningstar Portfolio Risk Score as text, e.g. 'Aggressive (73)': a single number for the current, forward-looking risk of the fund, based on the estimated volatility of its future returns (from its holdings). Usually 0-100; above 100 = extreme risk.",
                             "Morningstar bands: Conservative 0-23, Moderate 24-47, Aggressive 48-78, Very Aggressive 79-99, Extreme 100+. (Labels seen in this export use slightly different lower cutoffs: Conservative 1-28, Moderate 29-54, Aggressive 55-78.)",
                             "Lower for low-risk investing",
                             "Split into 'Risk Label' and 'Risk Score Numeric'.", MS, "No")
D["Beta"] = ("How strongly the fund moves with its benchmark: 1.0 = same swings, 1.2 = 20% bigger, 0.5 = half.",
             "Beta = Cov(fund excess return, benchmark excess return) / Var(benchmark excess return), monthly over 3Y",
             "Below 1 for lower risk; about 1 for index-like behaviour",
             "Only meaningful when the benchmark fits the fund. Negative = moves opposite (inverse / hedged).", MS, "No")
D["Alpha"] = ("Annualized return above (or below) what the fund's beta to its benchmark predicts, in percent.",
              "Alpha = (Rp - Rf) - Beta x (Rb - Rf), annualized; Rb = benchmark return",
              "Higher is better (> 0 = added value)",
              "Misleading if the benchmark doesn't fit the fund.", MS, "No")
D["Net Expense Ratio"] = ("Annual fee charged by the fund after fee waivers, in percent of assets (0.58 = 0.58% a year).",
                          "Net Expense Ratio = total annual operating expenses (after waivers) / average net assets",
                          "Lower is better",
                          "Waivers can expire. Index ETFs are often < 0.10%, active funds 0.5-1%. "
                          "Sector-neutral: driven by passive vs active, not holdings. Expense counts in the score through the Costs concept (Adjusted Expense Ratio).", MS, "No (Costs uses Adjusted Expense Ratio)")
D["Adjusted Expense Ratio"] = ("Morningstar's expense ratio excluding some costs that aren't true management fees (e.g. interest / dividend expense on short sales) - the 'real' ongoing cost.",
                               "Reported by Morningstar; not derived in this project", "Lower is better",
                               "Costs score = inverted within-category rank (cheapest in category = 100).", MS, used("costs", [("", "adjusted_expense_ratio")]))
D["Total Net Assets for Share Class"] = ("Assets in this share class only, in US dollars.", "Reported directly", "Higher is better (more stable)",
                                         "Under about $50-100M a fund has higher risk of closing. Use 'Fund Size' for the whole fund.", MS, "No")
D["Fund Size"] = ("Total assets of the whole fund (all share classes), in US dollars.",
                  "Liquidity/Size score (fixed cutoffs, not within category): ≤ $10M → 0, $500M → target, ≥ $5B → 100, linear in between",
                  "Higher is better up to a point",
                  f"Popularity / age driven; little sector link. 'Trading Volume' (YAML weight {CW['liquidity_size']['trading_volume']:g}) is not in the data, so Fund Size is effectively the whole Liquidity/Size score.",
                  MS, used("liquidity_size", [("", "fund_size")]))
D["Premium/Discount"] = ("ETFs / closed-end funds: how far the market price is above (+) or below (-) the value of the holdings (NAV), in percent. '1Y Avg' = average over the last year.",
                         "Premium/Discount = (Market Price - NAV) / NAV x 100", "Close to 0",
                         "Big premiums / discounts mean buying above or selling below fair value. Blank for mutual funds. The 1Y Avg variant is empty in the current export.", MS, "No")
D["Portfolio Growth Grade"] = ("Morningstar letter grade (A-F) of how fast the companies held are growing (revenue, earnings, cash flow).",
                               "Morningstar proprietary; converted to a number (config.GRADE_TO_NUMERIC) before ranking", "A best, F worst",
                               "Equity funds only. HIGH sector distortion: software screens with higher growth grades than hardware/semis.", MS, used("quality_valuation", [("", "growth_grade")]))
D["Portfolio Financial Health Grade"] = ("Morningstar letter grade (A-F) of the balance-sheet strength (debt, solvency) of the companies held.",
                                         "Morningstar proprietary; converted to a number before ranking", "A best, F worst", "Equity funds only.", MS, used("quality_valuation", [("", "financial_health")]))
D["Portfolio Profitability Grade"] = ("Morningstar letter grade (A-F) of how profitable the companies held are (return on equity, margins).",
                                      "Morningstar proprietary; converted to a number before ranking", "A best, F worst", "Equity funds only.", MS, used("quality_valuation", [("", "profitability_grade")]))
D["Portfolio Economic Moat Coverage"] = ("Percent of the stock holdings Morningstar rates as having a Wide moat (durable advantage 20+ years), Narrow moat (10+ years) or None.",
                                         "Wide + Narrow + None ≈ the part of the portfolio Morningstar analysts cover", "Wide: higher is better; None: lower is better",
                                         "Only for funds whose holdings Morningstar covers; blank for most bond / small-cap / foreign funds. Only 'Wide' is scored.", MS, used("quality_valuation", [("Wide", "economic_moat_wide")]))
D["Portfolio Return on Invested Capital"] = ("Average ROIC of the companies held, in percent - how efficiently they turn capital into profit.",
                                             "ROIC = after-tax operating profit / invested capital (holding-weighted)", "Higher is better (> 15% strong)",
                                             "Extreme values come from holdings with tiny or negative capital - unreliable.", MS, "No")
D["Portfolio Price/Earnings"] = ("Holding-weighted price-to-earnings ratio of the stocks held.", "P/E = price / earnings per share", "Lower = cheaper (value), higher = pricier (growth)",
                                 "Not good/bad on its own; growth funds always have higher P/E. Compare within category.", MS, "No")
D["Portfolio Price/Book"] = ("Holding-weighted price-to-book ratio.", "P/B = price / book value per share", "Lower = cheaper", "Banks / value funds are naturally low, tech is high.", MS, "No")
D["Portfolio Price/Sales"] = ("Holding-weighted price-to-sales ratio.", "P/S = price / sales per share", "Lower = cheaper", "Compare within category.", MS, "No")
D["Portfolio Price/Free Cash Flow"] = ("Holding-weighted price-to-free-cash-flow ratio.", "P/FCF = price / free cash flow per share", "Lower = cheaper (positive values)",
                                       "Negative = holdings burn cash; extreme values are unreliable.", MS, "No")
D["Portfolio Price/Fair Value"] = ("Market price of the holdings divided by Morningstar analysts' fair value estimates. 1.0 = fair, 0.8 = 20% undervalued, 1.2 = 20% overvalued.",
                                   "Price/Fair Value = Market Price / Morningstar Fair Value Estimate", "Lower is better (< 1 = undervalued)",
                                   "Only for funds whose holdings Morningstar covers. Scored inverted. HIGH sector distortion: software trades at richer valuations than hardware/semis.",
                                   MS, used("quality_valuation", [("", "price_fair_value")]))
D["ETF Fair Value"] = ("Morningstar's estimated fair value per ETF share, in dollars.", "", "Above the market price is better",
                       "Compare with the current price (not in the report); Price/Fair Value already gives the ratio.", MS, "No")
D["Tracking Error"] = ("How closely the fund follows its benchmark: annualized standard deviation of (fund return - benchmark return), in percent.",
                       "TE = sqrt( Σ[(Rp,i - Rb,i) - mean(Rp - Rb)]² / (n - 1) ) x sqrt(12)",
                       "Lower for index funds; depends for active funds",
                       "Low = hugs the benchmark. Active funds need some tracking error to beat the index. If the named benchmark has a different sector mix than the fund, TE is high for that reason, not poor management.",
                       MS, used("tracking", [("3Y", "tracking_error_3y"), ("1Y", "tracking_error_1y")]))
D["Fund Managers"] = ("Names of the current portfolio managers.", "", "n/a",
                      "Loaded but not scored. Check tenure on Morningstar - a recent manager change makes the past record less relevant.", MS, "No")
D["Medalist Rating"] = ("Morningstar's forward-looking rating. Overall: Gold / Silver / Bronze / Neutral / Negative. Pillars (Parent = fund company, People = managers, Process = strategy): High / Above Average / Average / Below Average / Low.",
                        "Overall rating combines the People, Process and Parent pillars with fees, relative to the category; converted to a number before ranking",
                        "Gold best ... Negative worst; pillars: High best",
                        "A view of FUTURE performance, unlike star ratings. Only 'Overall' is scored.", MS, used("quality_valuation", [("Overall", "medalist")]))
D["Morningstar Rating for Funds"] = ("The 1-5 star rating: past risk-adjusted return after fees and loads vs the category (Overall blends 3Y/5Y/10Y).",
                                     "Ranks Morningstar Risk-Adjusted Return within category: top 10% = 5 stars, next 22.5% = 4, middle 35% = 3, next 22.5% = 2, bottom 10% = 1",
                                     "5 best, 1 worst",
                                     "Backward-looking; needs 3+ years of history. Only 'Overall' is scored.", MS, used("quality_valuation", [("Overall", "star_rating")]))
D["Morningstar Return Rating"] = ("Morningstar's return rating vs the category: Low, Below Average, Average, Above Average, High.",
                                  "Same 10 / 22.5 / 35 / 22.5 / 10% bands as the star rating, on load-adjusted return only (no risk adjustment)",
                                  "High / Above Average", "Relative to the category.", MS, "No")
D["Tax Cost Ratio"] = ("Percent of return lost each year to taxes on the fund's distributions (highest tax bracket assumed).",
                       "Tax Cost Ratio = [1 - (1 + after-tax return) / (1 + pre-tax return)] x 100",
                       "Lower is better (taxable accounts)",
                       "Irrelevant in IRAs / 401(k)s. ETFs are usually more tax-efficient than mutual funds. Moderate sector link (mature dividend sectors distribute more).",
                       MS, used("tax_income", [("1Y", "tax_cost_ratio_1y"), ("2Y", "tax_cost_ratio_2y")]))
D["Potential Capital Gains Exposure"] = ("Percent of the fund's assets that are unrealized (or undistributed) gains - future taxable distributions in a taxable account.",
                                         "PCGE = unrealized appreciation / total net assets x 100", "Lower is better (taxable accounts)",
                                         "Irrelevant in IRAs / 401(k)s. Negative = losses that can offset future gains.", MS, "No")
D["SEC 30-Day Yield"] = ("Standardized yield from the last 30 days of net investment income, after expenses, annualized, in percent.",
                         "SEC Yield = 2 x [ ((a - b) / (c x d) + 1)^6 - 1 ]; a = income, b = expenses, c = avg shares, d = max offering price",
                         "Higher = more income",
                         "Fairest yield comparison between bond funds; usually blank for stock funds.", MS, used("tax_income", [("", "sec_yield")]))
D["Dividend per Share"] = ("Dollars of dividends per share: Trailing Annual = paid in the last 12 months; Forward Annual = expected next 12 months; 1Y / 3Y Avg = averages.",
                           "", "Higher = more income",
                           "Depends on the share price - compare yields rather than dollars across funds. Forward / 1Y Avg / 3Y Avg are empty in the current export.",
                           MS, used("tax_income", [("Trailing Annual", "dividend_per_share_trailing"), ("Forward Annual", "dividend_per_share_forward"),
                                                    ("1Y Avg", "dividend_per_share_1y_avg"), ("3Y Avg", "dividend_per_share_3y_avg")]))
D["Leveraged Fund"] = ("Yes if the fund uses borrowing / derivatives to multiply daily returns (2x, 3x, -1x).", "", "No for long-term holding",
                       "Leveraged funds reset daily and lose value in choppy markets ('volatility decay'). "
                       "Kept in every profile (no leveraged filter); compare them with care.", MS, "No")
D["Interval Fund"] = ("Yes if the fund only lets you sell at scheduled intervals (e.g. quarterly).", "", "No (more liquid)", "Kept and scored like other funds; check liquidity before buying.", MS, "No")
D["Fund of Funds"] = ("Yes if the fund invests in other funds (e.g. target-date funds).", "", "Neutral", "Check the total cost including the underlying funds' fees.", MS, "No")
D["Investment Status"] = ("Open, Limited (closed to most new investors) or Closed.", "", "Open", "You may not be able to buy a Closed / Limited fund.", MS, "No")
D["Strategic Beta Group"] = ("For 'smart beta' index funds: the factor they tilt to (Value, Dividend, Multi-Factor, ...).", "", "n/a", "'Not Applicable' or blank for regular funds.", MS, "No")
D["Share Class Type"] = ("The sales channel / fee structure of the share class: No Load, Front Load (A), Level Load (C), Deferred Load (B), Institutional, Retirement, Investor, Advisor; this export also uses 'ETF'.", "", "ETF / No Load / Institutional are usually cheapest",
                         "Retirement and Institutional classes may only be available through employer plans or at high minimums.", MS, "No")
D["Tender Offer"] = ("Yes if the fund only buys back shares through periodic tender offers (limited liquidity).", "", "No", "Kept and scored like other funds; check liquidity before buying.", MS, "No")
D["Index Fund"] = ("Yes if the fund tracks an index (passive); No = actively managed.", "", "Neutral (index funds are usually cheaper)", "", MS, "No")
D["Maximum Deferred Load"] = ("Maximum back-end sales charge (CDSC) when selling, in percent.", "", "0",
                              f"Purchase rule for every profile ({rule('max_deferred_load', pct)}, config.PURCHASE_RULES): funds above it are removed before scoring.", MS, "No (hard filter)")
D["Maximum Front Load"] = ("Maximum sales charge when buying, in percent.", "", "0", f"Purchase rule for every profile ({rule('max_front_load', pct)}, config.PURCHASE_RULES): funds above it are removed before scoring.", MS, "No (hard filter)")
D["Minimum Initial Investment"] = ("Minimum dollars needed for the first purchase.", "", "Lower is better (must fit your budget)",
                                   f"Purchase rule for every profile ({rule('min_initial_investment', dollars)}, config.PURCHASE_RULES): funds whose minimum is above it are removed. Blank = no minimum reported (ETFs: one share). Brokers may waive minimums.", MS, "No (hard filter)")
D["Redemption Fee"] = ("Fee charged when selling within a short holding period, in percent.", "", "0", f"Purchase rule for every profile ({rule('max_redemption_fee', pct)}, config.PURCHASE_RULES): funds above it are removed before scoring.", MS, "No (hard filter)")
D["Risk Label"] = ("Label part of 'Portfolio Risk Score': Conservative, Moderate, Aggressive, Very Aggressive or Extreme Risk.", "Text before the '(' in Portfolio Risk Score", "Lower risk for low-risk investing", "", CALC, "No")
D["Risk Score Numeric"] = ("Number part of 'Portfolio Risk Score' (about 100 = as risky as a diversified all-stock portfolio).", "Number inside the '( )' in Portfolio Risk Score", "Lower = less risky",
                           "Above 100 ('Extreme Risk') = riskier than an all-stock portfolio, e.g. leveraged, single-sector or crypto funds (361 funds in this report).", CALC, "No")
D["Yahoo Subsector"] = ("Industry group with the most weight among the fund's top 10 holdings (Yahoo), e.g. Semiconductors, Software & Services, Banks. 'Broad Market Equity' when no industry dominates or no holdings data.",
                        "For each of the top 10 holdings: look up its industry → sub-sector; sum holding weights per sub-sector; pick the largest",
                        "n/a (grouping)",
                        "Refreshed about once a year. Diversified funds near a tie (e.g. Nvidia vs Microsoft) can flip between neighbouring labels. Not meaningful for bond / allocation funds. Fills the sector gap of Morningstar Category; used as the peer group for Z-scores.", YH, "Used as peer group for Z Score")
D["Sharpe"] = ("Sharpe ratio computed by the screener from Yahoo daily prices over the last 12 (1Y) or 36 (3Y) months - an independent check of Morningstar's.",
               "Monthly return r = month-end Adj Close change; Sharpe = mean(r - 4%/12) / std(r - 4%/12) x sqrt(12). Needs ≥ 11 months (1Y) or ≥ 30 months (3Y)",
               "Higher is better (> 1 good, < 0 = worse than cash)",
               "Differences from Morningstar's Sharpe come from price data and the fixed 4% risk-free rate.", YH, used("risk_adjusted", [("3Y", "yahoo_sharpe_3y"), ("1Y", "yahoo_sharpe_1y")]))
D["Z Score"] = ("How many standard deviations the fund's Yahoo Sharpe is above (+) or below (-) the average of all funds with the same Yahoo Subsector.",
                "Z = (fund Sharpe - mean Sharpe of its sub-sector) / std of Sharpe in its sub-sector; 0 if the sub-sector has no spread",
                "Higher is better (> +1 clearly above peers, < -1 clearly below)",
                "Peer group = Yahoo Subsector (not Morningstar Category), so a bond fund labelled 'Broad Market Equity' is compared with stock funds. "
                "Read 1Y and 3Y together: both >= +1 = Core (long-term leader still leading); 1Y >= +1 but 3Y < 0 = Momentum (recent surge, tactical trade, not an anchor); "
                "1Y < 0 but 3Y >= +1 = Value (strong record in a temporary slump, possible buy-the-dip). List them with: python db_access/query.py zscore_matrix", CALC,
                used("risk_adjusted", [("3Y", "yahoo_zscore_3y"), ("1Y", "yahoo_zscore_1y")]))
D["Yahoo Fetch Error"] = ("Message when Yahoo returned no data for the fund.", "", "Blank", "Funds with no Yahoo prices are moved to the 'No Yahoo Data' tab.", CALC, "No")
D["Yahoo Data Suspect"] = ("1 if the fund is older than 3 years (per Morningstar) but has no Yahoo 3Y Sharpe / Z-score.", "Inception Date ≤ today - 3 years AND (Sharpe 3Y or Z Score 3Y blank)", "0",
                           "Usually an ETF converted from an older mutual fund (Yahoo only has the new ticker's prices).", CALC, "No")

# Label of each concept input, and whether "lower is better" (scored inverted)
INPUT_LABELS = {
    "long_term_return_performance": {f"return_{p}": (f"Total Return {p.upper()}", False)
                                     for p in ("1y", "2y", "3y", "4y", "5y", "6y", "7y", "8y", "9y", "10y", "15y", "20y")},
    "short_term_return_performance": {f"return_{p}": (f"Total Return {p.upper()}", False)
                                      for p in ("1w", "1m", "2m", "qtd", "3m", "6m", "9m", "ytd")},
    "risk_adjusted": {
        **{f"sharpe_{p}": (f"Morningstar Sharpe {p.upper()}", False) for p in ("1y", "3y", "5y", "10y")},
        "upside": ("Upside Capture 3Y", False), "downside": ("Downside Capture 3Y", True),
        **{f"upside_{p}": (f"Upside Capture {p.upper()}", False) for p in ("1y", "5y", "10y")},
        **{f"downside_{p}": (f"Downside Capture {p.upper()}", True) for p in ("1y", "5y", "10y")},
        "yahoo_sharpe_1y": ("Yahoo Sharpe 1Y", False), "yahoo_sharpe_3y": ("Yahoo Sharpe 3Y", False),
        "yahoo_zscore_1y": ("Z Score 1Y", False), "yahoo_zscore_3y": ("Z Score 3Y", False),
    },
    "volatility": {**{f"stdev_{p}": (f"Standard Deviation {p.upper()}", True) for p in ("1y", "3y", "5y", "10y")},
                   **{f"drawdown_{p}": (f"Max Drawdown {p.upper()}", False) for p in ("1y", "3y", "5y", "10y")}},
    "resilience": {"yearly_beat_pct": ("Yearly Beat %", False),
                   "yearly_bottom_10_share": ("Yearly Bottom 10% Count / Yearly Periods With Data", True),
                   "volatile_beat_pct": ("Volatile Beat %", False),
                   "volatile_bottom_10_share": ("Volatile Bottom 10% Count / Volatile Periods With Data", True),
                   "worst_3m_return": ("Worst Three Month Return", False)},
    "tracking": {"tracking_error_3y": ("Tracking Error 3Y", True), "tracking_error_1y": ("Tracking Error 1Y", True)},
    "liquidity_size": {"fund_size": ("Fund Size with fixed cutoffs ($10M → 0, $500M target, $5B → 100)", False),
                       "trading_volume": ("Trading Volume", False)},
    "quality_valuation": {"growth_grade": ("Growth Grade", False), "financial_health": ("Financial Health", False),
                          "profitability_grade": ("Profitability", False), "price_fair_value": ("Price/Fair Value", True),
                          "medalist": ("Medalist", False), "star_rating": ("Star rating", False),
                          "economic_moat_wide": ("Wide moat", False)},
    "costs": {"adjusted_expense_ratio": ("Adjusted Expense Ratio", True)},
    "tax_income": {"tax_cost_ratio_1y": ("Tax Cost Ratio 1Y", True), "tax_cost_ratio_2y": ("Tax Cost Ratio 2Y", True),
                   "dividend_per_share_trailing": ("Dividend per Share (Trailing)", False),
                   "dividend_per_share_forward": ("Dividend per Share (Forward)", False),
                   "dividend_per_share_1y_avg": ("Dividend per Share (1Y Avg)", False),
                   "dividend_per_share_3y_avg": ("Dividend per Share (3Y Avg)", False),
                   "sec_yield": ("SEC 30-Day Yield", False)},
}


def inputs_text(concept):
    """'Total Return 5Y 15%, ...' for the inputs this YAML weights, as shares of the concept."""
    parts = []
    for key, (label, inverted) in INPUT_LABELS[concept].items():
        if CW.get(concept, {}).get(key) and (concept, key) not in ABSENT:
            parts.append(f"{label} {inner_share(concept, key):.0%}" + (" (inverted)" if inverted else ""))
    return ", ".join(parts) or "none (all input weights are 0)"


concept_rows = [
    ("Long Term Return Performance Score", "long_term_return_performance", "Missing periods (young funds) are skipped, not counted as 0."),
    ("Short Term Return Performance Score", "short_term_return_performance", "Recent momentum - changes quickly."),
    ("Risk Adjusted Score", "risk_adjusted", "Missing Yahoo values drop out instead of counting as 0."),
    ("Volatility Score", "volatility", "Blank for funds younger than 3 years. Max drawdown is negative, so it is NOT inverted: closer to 0 scores higher."),
    ("Resilience Score", "resilience", "How consistently the fund beat the indexes: calendar years and the S&P 500's volatile periods (crashes and recoveries). Bottom-10% counts are divided by the periods the fund has data for, so older funds aren't penalized for having more periods."),
    ("Tracking Score", "tracking", "Good for index funds; for active funds low tracking error just means 'close to the index'."),
    ("Liquidity Size Score", "liquidity_size", "Not within category."),
    ("Quality Valuation Score", "quality_valuation", "Mostly equity data; bond funds rely on Medalist and stars."),
    ("Costs Score", "costs", ""),
    ("Tax Income Score", "tax_income", ""),
]
for name, concept, note in concept_rows:
    cw = W.get(concept, 0)
    D[name] = (f"0-100 concept score from these inputs (YAML weights, divided by their total): {inputs_text(concept)}.", NORM, "Higher is better (100 = best in category)",
               (WITHIN + " " + note).strip(), CALC, f"Yes - YAML weight {cw:g} = {concept_share(concept):.1%} of {PA}" if cw else f"No - 0 weight in {PA}")
D["Return Rank"] = ("Rank of the matching 'Total Return' column across ALL funds in the report (1 = highest return).",
                    "Rank = position when all funds are sorted by that period's return, highest first (ties share the lowest rank)",
                    "Lower is better (1 = best)",
                    "Across all categories, not within category - leveraged / sector funds dominate both ends. For momentum lists use python db_access/query.py momentum_top / momentum_bottom.", CALC, "No")
def add_profile_rows(key, yaml_name, weights):
    """The score and rank columns one profile adds to the report."""
    pa, total = f"Profile {key}", sum(v for v in weights.values() if v)
    D[f"{pa} Score"] = (f"Final 0-100 score combining the concept scores with the weights in {yaml_name}.",
                        f"{pa} Score = Σ(concept weight x concept score) / Σ(weights of concept scores the fund has). YAML weights: " + ", ".join(f"{CONCEPT_LABEL[k]} {v:g}" for k, v in weights.items()) + f" (sum {total:g}, so each effective share = weight / {total:g})",
                        "Higher is better",
                        "Blank concept scores (e.g. Volatility for young funds) are left out and the rest re-weighted. Because inputs are ranked within category, the score means 'good for its category', not 'better than every other fund'."
                        + ("" if key == KEY else f" Uses {key}'s own concept weights, so it can't be recomputed from the concept score columns in this report (those use {KEY}'s). Blank for funds {yaml_name} filters out."),
                        CALC, f"This IS the {pa} score" if key == KEY else f"No - this is the {pa} score")
    D[f"{pa} Rank In Category"] = (f"Rank of the {pa} Score within the fund's Morningstar Category (1 = best).", "", "Lower is better",
                                   "Read with the category size: 5th of 8 is mediocre, 5th of 400 is excellent. The highest numbers are the category's weakest funds for this profile.", CALC, "No")
    D[f"{pa} Rank Overall"] = (f"Rank of the {pa} Score across all funds (1 = best).", "", "Lower is better",
                               "Mixes categories; a top bond fund and a top stock fund can both rank high.", CALC, "No")


add_profile_rows(KEY, PROFILE_PATH.name, W)
# A combined report (results_profile_LT_ST_...) also has the other profiles' columns.
for other in ALL_PROFILES:
    if other.profile_name != KEY and f"Profile {other.profile_name} Score" in report_cols:
        add_profile_rows(other.profile_name, Path(other.path).name, other.thresholds["weights"])
D["Missing From Yahoo"] = ("('No Yahoo Data' tab only) Which Yahoo-based results are missing: Sharpe / Z-score, yearly returns, volatile periods.", "", "n/a",
                           "These funds are taken off the overview; all their Morningstar columns are still on that tab.", CALC, "No")

# Columns the YAML weights mention that are not in the current data
NOT_IN_DATA = "Morningstar (not in current data)"
D["Trading Volume"] = ("Average number of shares traded per day (ETFs).", "Reported directly (exchange data)", "Higher = easier to trade at a fair price",
                       f"Not in the current eTrade export. The YAML still gives it weight {CW['liquidity_size']['trading_volume']:g} in Liquidity/Size; since the column is missing, Fund Size gets the whole Liquidity/Size weight. Add it to the export to use it.",
                       NOT_IN_DATA, f"No effect: weighted {CW['liquidity_size']['trading_volume']:g} in the YAML, but the column is missing")
groups.setdefault("Trading Volume", [])

# Per-column score usage for the All Report Columns tab: {base: {variant: (concept, key)}}
VARIANT_KEYS = {
    "Total Return": {**{p: ("short_term_return_performance", f"return_{p.lower()}") for p in ["1W", "1M", "2M", "QTD", "3M", "6M", "9M", "YTD"]},
                     **{p: ("long_term_return_performance", f"return_{p.lower()}") for p in ["1Y", "2Y", "3Y", "4Y", "5Y", "6Y", "7Y", "8Y", "9Y", "10Y", "15Y", "20Y"]}},
    "Sharpe Ratio": {"1Y": ("risk_adjusted", "sharpe_1y"), "3Y": ("risk_adjusted", "sharpe_3y"),
                     "5Y": ("risk_adjusted", "sharpe_5y"), "10Y": ("risk_adjusted", "sharpe_10y")},
    "Upside Capture Ratio": {"1Y": ("risk_adjusted", "upside_1y"), "3Y": ("risk_adjusted", "upside"),
                             "5Y": ("risk_adjusted", "upside_5y"), "10Y": ("risk_adjusted", "upside_10y")},
    "Downside Capture Ratio": {"1Y": ("risk_adjusted", "downside_1y"), "3Y": ("risk_adjusted", "downside"),
                               "5Y": ("risk_adjusted", "downside_5y"), "10Y": ("risk_adjusted", "downside_10y")},
    "Sharpe": {"3Y": ("risk_adjusted", "yahoo_sharpe_3y"), "1Y": ("risk_adjusted", "yahoo_sharpe_1y")},
    "Z Score": {"3Y": ("risk_adjusted", "yahoo_zscore_3y"), "1Y": ("risk_adjusted", "yahoo_zscore_1y")},
    "Standard Deviation": {p: ("volatility", f"stdev_{p.lower()}") for p in ("1Y", "3Y", "5Y", "10Y")},
    "Maximum Drawdown": {p: ("volatility", f"drawdown_{p.lower()}") for p in ("1Y", "3Y", "5Y", "10Y")},
    "Tracking Error": {"3Y": ("tracking", "tracking_error_3y"), "1Y": ("tracking", "tracking_error_1y")},
    "Tax Cost Ratio": {"1Y": ("tax_income", "tax_cost_ratio_1y"), "2Y": ("tax_income", "tax_cost_ratio_2y")},
    "Medalist Rating": {"Overall": ("quality_valuation", "medalist")},
    "Morningstar Rating for Funds": {"Overall": ("quality_valuation", "star_rating")},
    "Portfolio Economic Moat Coverage": {"Wide": ("quality_valuation", "economic_moat_wide")},
    "Dividend per Share": {"Trailing Annual": ("tax_income", "dividend_per_share_trailing")},
}


def tab_of(col):
    if col == "Missing From Yahoo":
        return "No Yahoo Data"
    return "Overview"


all_rows = []
for col in report_cols + extra_cols:
    base, variant = base_of(col)
    d, f, o, c, s, u = D[base]
    if base in VARIANT_KEYS:
        ck = VARIANT_KEYS[base].get(variant)
        u = used(ck[0], [("", ck[1])]) if ck else "No"
    all_rows.append({"Report Column": col, "Tab": tab_of(col), "Column Dictionary Row": base,
                     "Variant": variant, "Optimal Value": o, "Source": s, f"Used In {PA} Score": u})
all_columns = pd.DataFrame(all_rows)

# Morningstar's own definitions (docs/morningstar_definitions.json, researched
# from Morningstar's glossary and methodology documents). Where ours was wrong
# ("corrects"), Morningstar's definition and formula replace it.
MS_CHECK = {"confirms": "Matches Morningstar", "corrects": "Corrected to Morningstar's definition",
            "not found": "No Morningstar definition found"}
MS_DEFS = {r["column"]: r for r in json.load(open(PROJECT / "docs/morningstar_definitions.json", encoding="utf-8"))}
KEEP_OURS = {"Portfolio Risk Score", "Share Class Type"}  # corrected by hand above (data differs from the glossary)
for name, r in MS_DEFS.items():
    if name in D and r["verdict"] == "corrects" and name not in KEEP_OURS:
        d, f, o, c, s_, u = D[name]
        D[name] = (r["morningstar_definition"], r["formula"] or f, o, c, s_, u)


def ms_columns(base):
    r = MS_DEFS.get(base)
    if not r:
        return {"Morningstar Definition": "", "Morningstar Check": "", "Morningstar Source": ""}
    extra = f" Key facts: {r['key_facts']}" if r.get("key_facts") else ""
    source = r["source_url"] + ("" if r["source_type"].startswith("Morningstar") else f" ({r['source_type']})")
    return {"Morningstar Definition": (r["morningstar_definition"] + extra).strip(),
            "Morningstar Check": MS_CHECK.get(r["verdict"], r["verdict"]), "Morningstar Source": source}


# --------------------------------------------------------------------------
rows = []
for base, variants in groups.items():
    if base not in D:
        raise KeyError(f"no definition for {base!r}")
    d, f, o, c, s, u = D[base]
    if variants == ["current"]:
        variants = []
    if base in ("Volatile Decline", "Volatile Recovery"):
        shown = "; ".join(variants)
    else:
        shown = ", ".join(variants)
    rows.append({"Column Name": base, "Variants": shown, "Definition": d, "Formula / Notes": f,
                 "Optimal Value": o, "Condition": c, "Source": s, f"Used In {PA} Score": u, **ms_columns(base)})
dictionary = pd.DataFrame(rows)

# Volatile periods tab
vp = pd.DataFrame([{
    "Period": p["name"], "Type": p["type"].replace("_", " "), "Start (S&P high)": p["start"], "Low": p["end"],
    "Recovered old high": p.get("recovery", ""), "S&P 500 change %": p["sp500_change_pct"],
    "Peak 21-day volatility %": p["peak_21d_vol_pct"], "Trading days": p["trading_days"],
} for p in PERIODS.values()])

# Concept tab (weights from the current YAML)
concepts = pd.DataFrame([
    ("Long-term Return", W["long_term_return_performance"], "Total Return 1Y-20Y", "Medium-High", "Yes / Medium-High", "Multi-year returns diverge a lot by sub-sector timing."),
    ("Short-term Return", W["short_term_return_performance"], "Total Return 1W-YTD", "Medium-High", "Yes / Medium-High", "Short windows are dominated by what's in favour this month."),
    ("Risk-adjusted", W["risk_adjusted"], "Sharpe, capture ratios, Yahoo Sharpe & Z-scores", "Medium", "Yes / Moderate", "Inherits volatility's distortion through the Sharpe denominator."),
    ("Volatility", W["volatility"], "Standard Deviation and Max Drawdown 1Y/3Y/5Y/10Y", "High", "Yes / High", "Std dev and drawdown directly reflect sub-sector cyclicality (semis >> software)."),
    ("Resilience", W.get("resilience", 0), "Yearly / Volatile Beat %, share of years / periods in the bottom 10%, Worst Three Month Return", "Medium", "No / Moderate", "Beating ALL the broad indexes favours whatever led the market (e.g. tech in 2023-24); defensive funds shine in the crash periods instead."),
    ("Tracking", W["tracking"], "Tracking Error 1Y/3Y", "Low-Medium", "Yes / Low", "Benchmark-relative, so partly self-correcting."),
    ("Liquidity/Size", W["liquidity_size"], "Fund Size", "Low", "No / No", "Driven by popularity and age, not holdings."),
    ("Quality/Valuation", W["quality_valuation"], "Grades, Price/Fair Value, Medalist, stars, moat", "High", "Yes / High", "Growth grade and Price/Fair Value differ systematically software vs hardware."),
    ("Costs", W["costs"], "Adjusted Expense Ratio", "Low", "No / No", "Driven by passive vs active, not composition."),
    ("Tax/Income", W["tax_income"], "Tax Cost Ratio, dividends, SEC yield", "Medium", "Yes / Moderate", "Yield and distributions differ by sector maturity."),
], columns=["Concept", "YAML weight", "Inputs", "Sector-composition distortion risk",
            "Depends on Morningstar Category / on dominant industry", "Why"])
concepts.insert(2, f"Effective share of {PA}", concepts["YAML weight"].map(lambda v: f"{v / TOTAL_CW:.1%}"))

# Sector reference: the 11 GICS-style sectors and the sub-sectors that behave differently inside them
sectors = pd.DataFrame([
    ("Technology", "Growth anchor", "Highest internal dispersion; separating sub-sectors matters most here.",
     "Semiconductors & Equipment (NVDA, AMD, TSM); Software & Services (MSFT, ORCL, CRM); Hardware, Storage & Peripherals (AAPL, HPQ); IT Services & Consulting (ACN, INFY)"),
    ("Healthcare", "Defensive growth", "Defensive stability mixed with volatile, speculative biotech research.",
     "Pharmaceuticals (LLY, JNJ, PFE); Biotechnology (AMGN, VRTX); Equipment & Supplies (MDT, ABT); Providers & Services (UNH, ELV)"),
    ("Financials", "Interest-rate cyclical", "Sensitive to central-bank rates, credit expansion and the yield curve.",
     "Banks (JPM, BAC, WFC); Financial Services / Capital Markets (V, MA, MS, GS); Insurance (BRK.B, PGR, MET)"),
    ("Communication Services", "Modern media", "Classic telecom combined with digital advertising and entertainment.",
     "Interactive Media (GOOGL, META); Entertainment & Streaming (NFLX, DIS); Telecom Services (T, VZ)"),
    ("Consumer Cyclical", "Economic pulse", "Relies on disposable income: thrives in expansions, contracts sharply in recessions.",
     "Automobiles (TSLA, F, GM); Broadline Retail & E-Commerce (AMZN); Hotels, Restaurants & Leisure (MCD, SBUX, BKNG); Durables & Apparel (NKE, HD)"),
    ("Consumer Defensive", "Recession guard", "Essentials bought in any economy: low volatility, steady dividends.",
     "Beverages & Tobacco (KO, PEP, PM); Food & Staples Retailing (WMT, COST); Household & Personal Products (PG, CL)"),
    ("Energy", "Commodity cycle", "Tied to oil and gas prices and alternative power.",
     "Oil, Gas & Fuels (XOM, CVX); Equipment & Services (SLB, HAL)"),
    ("Industrials", "Capital expenditure", "Machinery, manufacturing, construction and logistics.",
     "Aerospace & Defense (LMT, RTX); Transportation & Logistics (UPS, FDX, UNP); Machinery & Electrical Equipment (CAT, GE)"),
    ("Materials", "Raw base", "Front of the supply chain: turns raw elements into materials for manufacturers.",
     "Chemicals (LIN, DD); Metals & Mining (FCX, NUE, NEM); Containers & Packaging (IP)"),
    ("Utilities", "Regulated income bond", "Heavily regulated and debt-financed, so they trade like a bond proxy.",
     "Electric, Gas & Water (NEE, DUK, SO); Independent Power & Renewables (CEG)"),
    ("Real Estate", "Hard-asset yield", "Owners and managers of income-producing property.",
     "Equity REITs (PLD, AMT, O); Management & Development (CBRE)"),
], columns=["Sector", "Character", "Behaviour", "Sub-sectors (example tickers)"])

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    dictionary.to_excel(xw, index=False, sheet_name="Column Dictionary")
    all_columns.to_excel(xw, index=False, sheet_name="All Report Columns")
    concepts.to_excel(xw, index=False, sheet_name="Scoring Concepts")
    vp.to_excel(xw, index=False, sheet_name="Volatile Periods")
    sectors.to_excel(xw, index=False, sheet_name="Sectors")

wb = load_workbook(OUT)
fill = PatternFill("solid", fgColor="DDEBF7")
widths = {
    "Column Dictionary": [30, 28, 55, 50, 28, 60, 18, 45, 60, 22, 45],
    "All Report Columns": [55, 28, 32, 30, 32, 22, 60],
    "Scoring Concepts": [20, 10, 14, 40, 18, 24, 60],
    "Volatile Periods": [55, 16, 14, 12, 16, 12, 14, 10],
    "Sectors": [24, 22, 55, 110],
}
for ws in wb.worksheets:
    for i, w in enumerate(widths[ws.title]):
        ws.column_dimensions[chr(65 + i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

# Notes under the concept table on why within-category ranking has limits
SECTOR_NOTES = [
    "Why sector mix matters: every score is ranked WITHIN Morningstar Category, but categories are coarse. "
    "Morningstar's 'Technology' bundles software (operating systems, applications) with hardware, storage, networking "
    "and semiconductors. Two funds in the same category can be semiconductor-heavy (cyclical, capital-intensive) or "
    "software-heavy (higher margins, recurring revenue) and still be ranked as direct peers on every score.",
    "Practical takeaway: before trusting a ranking or tuning the Volatility, Quality/Valuation or Return weights in a "
    "sector-heavy category like Technology, check the Yahoo Subsector column (or the fund's factsheet holdings) "
    "of the top funds. Costs and Liquidity/Size are the least affected; Volatility and Quality/Valuation the most.",
]
ws = wb["Scoring Concepts"]
for offset, note in enumerate(SECTOR_NOTES):
    row = ws.max_row + 2 if offset == 0 else ws.max_row + 1
    ws.cell(row=row, column=1, value=note).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    ws.row_dimensions[row].height = 75
wb.save(OUT)
print(f"{len(dictionary)} column names, {len(all_columns)} report columns -> {OUT}")
print(dictionary["Source"].value_counts().to_string())
